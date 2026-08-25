"""Bóc chữ từ tệp khách gửi — dựng tệp THẬT rồi đọc lại, không giả lập."""

import io

import pytest

from app.services import documents as doc


def _xlsx(rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Báo giá"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx(paras, table=None):
    import docx

    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.rows[i].cells[j].text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(lines):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def test_doc_excel_giu_duoc_so_lieu():
    data = _xlsx([["Mặt hàng", "SL", "Đơn giá"], ["Gioăng đai treo", 100, 25000]])
    ok, text = doc.extract("baogia.xlsx", None, data)
    assert ok
    assert "Gioăng đai treo" in text and "25000" in text
    assert "Sheet: Báo giá" in text


def test_doc_word_doc_ca_bang():
    data = _docx(["Yêu cầu báo giá"], table=[["Quy cách", "Số lượng"], ["Phi 21", "200"]])
    ok, text = doc.extract("yeucau.docx", None, data)
    assert ok
    assert "Yêu cầu báo giá" in text
    assert "Phi 21" in text and "200" in text  # bảng trong Word hay chứa đúng phần cần đọc


def test_doc_csv():
    ok, text = doc.extract("ds.csv", "text/csv", "tên,sl\ngioăng,10\n".encode())
    assert ok and "gioăng | 10" in text


def test_doc_pdf_scan_khong_co_chu_bao_ro():
    ok, text = doc.extract("scan.pdf", "application/pdf", _pdf([]))
    assert ok
    assert "không có chữ đọc được" in text  # nói thẳng thay vì trả chuỗi rỗng khó hiểu


def test_doc_dinh_dang_khong_ho_tro():
    ok, text = doc.extract("video.mp4", "video/mp4", b"\x00\x01")
    assert not ok and "chưa hỗ trợ" in text


def test_doc_qua_nang():
    ok, text = doc.extract("to.pdf", "application/pdf", b"x" * (doc.MAX_BYTES + 1))
    assert not ok and "nặng quá" in text


def test_doc_rong():
    ok, text = doc.extract("a.pdf", "application/pdf", b"")
    assert not ok and "rỗng" in text


def test_doc_hong_khong_lam_sap():
    ok, text = doc.extract("hong.xlsx", None, b"day khong phai excel")
    assert not ok and "mở tệp này không được" in text


def test_doc_cat_bot_khi_qua_dai():
    big = _xlsx([[f"dòng {i}", "x" * 200] for i in range(500)])
    ok, text = doc.extract("to.xlsx", None, big)
    assert ok
    assert len(text) <= doc.MAX_CHARS + 60
    assert "đã cắt" in text or "chỉ đọc" in text


@pytest.mark.parametrize("name,ct,mong", [
    ("a.PDF", None, "pdf"),
    ("b.xlsx", None, "xlsx"),
    ("c.docx", None, "docx"),
    ("d.txt", None, "text"),
    ("e.bin", None, None),
    ("f", "application/pdf", "pdf"),
    ("g", "text/csv; charset=utf-8", "csv"),  # trình duyệt hay gắn thêm charset
])
def test_doc_nhan_dien_loai(name, ct, mong):
    assert doc.kind_of(name, ct) == mong
